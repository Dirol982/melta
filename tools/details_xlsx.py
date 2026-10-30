"""Таблица характеристик товаров для заполнения и обратная загрузка в tools/details.json.

python3 tools/details_xlsx.py export docs/melta-harakteristiki.xlsx   — выгрузить (заполненное сохраняется)
python3 tools/details_xlsx.py import docs/melta-harakteristiki.xlsx   — загрузить заполненное
"""
import json, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
LIB = json.loads((ROOT / 'tools/library.json').read_text(encoding='utf-8'))
DETAILS_PATH = ROOT / 'tools/details.json'
DETAILS = json.loads(DETAILS_PATH.read_text(encoding='utf-8'))

# (ключ, заголовок, ширина, подсказка); ключи specs.* лежат в details[slug]['specs']
FIELDS = [
    ('volume', 'Объём / вес', 12, 'Как на упаковке: 500 мл, 1 кг, 100 г'),
    ('lead', 'Короткое описание', 40, '2–3 предложения: что это и чем хорош продукт'),
    ('benefits', 'Преимущества', 40, 'Каждое с новой строки: «Заголовок — пояснение». Только подтверждённые'),
    ('audience', 'Кому подходит', 28, 'Тип кожи/волос, возраст, вид животного, поверхности'),
    ('composition', 'Состав (INCI)', 40, 'Полностью, как на этикетке'),
    ('usage', 'Способ применения', 40, 'Каждый шаг с новой строки; меры предосторожности'),
    ('specs.weight', 'Вес нетто', 12, ''),
    ('specs.shelf', 'Срок годности', 14, 'Например: 24 месяца'),
    ('specs.storage', 'Условия хранения', 24, ''),
    ('specs.barcode', 'Штрихкод', 16, '13 цифр'),
    ('specs.box', 'Кратность в коробке', 12, 'Сколько штук в коробке'),
    ('specs.doc', 'Документ', 22, 'ТУ, декларация, СГР'),
    ('wb', 'Ссылка на Wildberries', 30, 'Ссылка на карточку товара'),
    ('back_photo', 'Фото оборота этикетки', 14, 'Есть / нет'),
]
LISTS = {'benefits', 'usage'}
FILL_TODO = PatternFill('solid', fgColor='FFF1E0')
FILL_HEAD = PatternFill('solid', fgColor='141019')


def get(slug, key, product):
    d = DETAILS.get(slug, {})
    if key == 'volume':
        return product['volume']
    if key.startswith('specs.'):
        return d.get('specs', {}).get(key[6:], '')
    v = d.get(key, '')
    if key == 'benefits':
        return '\n'.join(f"{b['title']} — {b['text']}" if isinstance(b, dict) and b.get('text') else (b['title'] if isinstance(b, dict) else b) for b in v)
    return '\n'.join(v) if isinstance(v, list) else v


def export(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Характеристики'
    head = ['Код', 'Категория', 'Товар', 'Уже есть (с этикетки)'] + [f[1] for f in FIELDS]
    ws.append(head)
    ws.append(['', '', '', 'Заполнено по фото — проверьте'] + [f[3] for f in FIELDS])
    for p in sorted(LIB['products'], key=lambda p: (list(LIB['categories']).index(p['cat']), p['name'])):
        d = DETAILS.get(p['slug'], {})
        known = '; '.join(d.get('label', []))
        if d.get('oils'):
            known += ('; ' if known else '') + 'Эфирные масла: ' + d['oils']
        ws.append([p['slug'], LIB['categories'][p['cat']], p['name'].split(': ')[0], known] + [get(p['slug'], f[0], p) for f in FIELDS])
    widths = [30, 18, 44, 40] + [f[2] for f in FIELDS]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for c in ws[2]:
        c.font = Font(italic=True, color='5B5568', size=9)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if c.column > 4 and not c.value:
                c.fill = FILL_TODO
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A1:{ws.cell(1, len(head)).column_letter}{ws.max_row}'
    dv = DataValidation(type='list', formula1='"Есть,Нет"', allow_blank=True)
    col = ws.cell(1, len(head)).column_letter
    dv.add(f'{col}3:{col}{ws.max_row}')
    ws.add_data_validation(dv)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print('выгружено', ws.max_row - 2, 'товаров →', path)


def import_(path):
    ws = load_workbook(path).active
    head = [c.value for c in ws[1]]
    lib = {p['slug']: p for p in LIB['products']}
    n = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        slug = row[0]
        if slug not in lib:
            continue
        d = DETAILS.setdefault(slug, {'label': [], 'specs': {}})
        d.setdefault('specs', {})
        for key, title, *_ in FIELDS:
            v = row[head.index(title)]
            v = '' if v is None else str(v).strip()
            if not v:
                continue
            if key == 'volume':
                lib[slug]['volume'] = v
            elif key.startswith('specs.'):
                d['specs'][key[6:]] = v
            elif key == 'benefits':
                d[key] = [dict(zip(('title', 'text'), [x.strip() for x in line.split(' — ', 1)])) for line in v.splitlines() if line.strip()]
            elif key in LISTS:
                d[key] = [line.strip() for line in v.splitlines() if line.strip()]
            else:
                d[key] = v
        n += 1
    DETAILS_PATH.write_text(json.dumps(DETAILS, ensure_ascii=False, indent=1), encoding='utf-8')
    (ROOT / 'tools/library.json').write_text(json.dumps(LIB, ensure_ascii=False, indent=2), encoding='utf-8')
    print('загружено строк:', n, '— теперь запустите tools/build_library.py и tools/build_catalog_js.py')


if __name__ == '__main__':
    {'export': export, 'import': import_}[sys.argv[1]](sys.argv[2])
